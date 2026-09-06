# -*- coding: utf-8 -*-
"""从与 HTML 报告相同的内容源生成 Word 版开发报告"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

PURPLE = RGBColor(0x6D, 0x5D, 0xF6)
DARK = RGBColor(0x2B, 0x1B, 0x6B)
GRAY = RGBColor(0x6B, 0x6B, 0x7B)

doc = Document()

# 全局中文字体
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

def set_cn(run, name="微软雅黑"):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)

def para(text, bold=False, color=None, size=None, align=None, indent=True):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if color: r.font.color.rgb = color
    if size: r.font.size = Pt(size)
    set_cn(r)
    if align: p.alignment = align
    if indent and not bold:
        p.paragraph_format.first_line_indent = Pt(21)
    return p

def h1(text):
    p = doc.add_heading(level=1)
    r = p.add_run(text)
    r.font.color.rgb = DARK
    r.font.size = Pt(16)
    set_cn(r)

def h2(text):
    p = doc.add_heading(level=2)
    r = p.add_run(text)
    r.font.color.rgb = PURPLE
    r.font.size = Pt(13)
    set_cn(r)

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for i, htxt in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        r = cell.paragraphs[0].add_run(htxt)
        r.bold = True
        set_cn(r)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.size = Pt(9.5)
            set_cn(r)
    return t

# ===== 封面 =====
for _ in range(6): doc.add_paragraph()
para("Orbit 日程助手 · 开发报告", bold=True, color=DARK, size=26, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
para("2026年9月6日 · 从零到真机可用 · 约14小时", color=GRAY, size=12, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
doc.add_paragraph()
para("环境：Windows（无Mac）· GitHub Actions 云构建 · Sideloadly 侧载 · iPhone 15 Pro（iOS 26.6.1）",
     color=GRAY, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
para("本 Word 版与 docs/开发报告.html 同源同内容；HTML 为主文档，本文件用于批注编辑。",
     color=GRAY, size=9, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
doc.add_page_break()

# ===== 1 项目概览 =====
h1("1 · 项目概览")
h2("1.1 需求")
para("iPhone 15 Pro（iOS 26）上的日程助手：支持语音 / 图片 / 文字三种输入，AI 自动识别为结构化日程，写入手机系统日历；支持多家大模型（保留扩展接口）。")
h2("1.2 成品（Orbit v2.1）功能清单")
table(["模块", "功能"], [
    ["聊天主界面", "消息流交互：用户气泡（文字/图片/语音转写）→ 助手回复 + 日程卡片"],
    ["输入方式", "底部输入栏：＋（照片/相机）｜文字框｜🎤 语音（点击录音、实时转写、自动发送）"],
    ["AI 识别", "多 LLM 统一接口：智谱 GLM（默认，多模态）/ OpenAI / 任意 OpenAI 兼容端点"],
    ["批量识别", "多行日程逐项识别；整场会议（名称+起止时间）单项录入"],
    ["日程卡片", "emoji+标题+日期+时间；提醒开关+时长（准时/5/15/30/60/120分钟/1天）；修改日程=切换日历；点卡片编辑/删除"],
    ["自动写日历", "识别成功自动写入默认日历（EventKit，兼容仅写入降级）"],
    ["日程页", "主页右上角：按时间排序列表，点击修改、左滑删除（同步删日历事件）"],
    ["侧边抽屉", "6套图标切换；API配置+测试；默认日历/提醒时长；使用说明"],
    ["数据", "聊天记录 JSON 持久化，重启可查"],
])
h2("1.3 技术栈")
para("SwiftUI（iOS 18）· EventKit · Speech · PhotosUI · URLSession async/await · GitHub Actions（macos-14，Xcode 16.2）双产物构建 · Sideloadly 免费证书侧载")
para("仓库：github.com/hanwen-FDE/orbit-scheduler")

# ===== 2 前置条件 =====
doc.add_page_break()
h1("2 · iOS 开发前置条件（先读）")
para("为什么这一章最重要：今天多数排障时间浪费在『没人提前讲清条件』上——开发者模式默认隐藏、免费 Apple ID 的密码规则、Windows 无法编译 iOS 等。以下条件应在动手前确认。", bold=True, color=RGBColor(0xB4,0x53,0x09))
table(["条件", "说明", "今天是否踩坑"], [
    ["代码编写", "任何平台（Windows 也行），写 Swift 源码 + Xcode 工程文件", "否"],
    ["编译", "只能在 macOS + Xcode。替代：GitHub Actions 云构建（本方案）/ 云Mac / 借Mac", "是：本机模拟器方案直接失败"],
    ["GitHub 账号", "公开仓库 macOS 云构建免费；私有按10倍分钟计费", "是：git 不走系统代理推送失败"],
    ["Apple ID", "免费 Apple ID 即可，无需688元开发者账号；签名7天过期需重签", "是：误以为需要开发者账号"],
    ["开发者模式", "iOS 16+ 侧载必须开启；开关默认隐藏，需 iMazing 触发显示", "是（最大坑）：Invalid file 的根因"],
    ["签名工具", "Sideloadly v0.60+（支持 iOS 26）；爱思助手证书环节失败弃用", "是"],
    ["App 权限描述", "Info.plist 须声明每个权限用途；漏一个=崩溃或静默失败", "是：日历静默失败、相机闪退"],
])
para("费用：全程 0 元。唯一代价是每 7 天连电脑重签一次。消除限制：Apple Developer Program 688元/年（可走 TestFlight，无需开发者模式）。")

# ===== 3 五种路径 =====
h1("3 · 无 Mac 的五种 iOS 开发路径")
table(["路径", "成本", "优点", "缺点", "适合"], [
    ["① GitHub Actions 云构建 + Sideloadly（本项目）", "0元", "全自动构建；日常只需浏览器+数据线", "7天重签；无法断点调试", "功能开发、自用"],
    ["② Appetize.io 网页模拟", "免费约100分钟/月", "浏览器直接跑 .app", "无麦克风/相机/推送", "快速验证UI"],
    ["③ 云 Mac（MacinCloud 等）", "约$1/小时起", "完整 Xcode 调试", "付费；有延迟", "认真持续开发"],
    ["④ 借/买 Mac", "人情/自购", "体验最佳", "依赖他人", "长期开发者"],
    ["⑤ 付费开发者 + TestFlight", "688元/年", "签名一年有效；TestFlight 免开发者模式；可分发", "年费；上传构建仍需 Mac（可用①的CI）", "正式/多人使用"],
])
para("组合建议：个人自用＝①+②；长期投入＝③或④；给同事装＝⑤。", bold=True)

# ===== 4 时间线 =====
doc.add_page_break()
h1("4 · 今日开发时间线")
table(["阶段", "内容"], [
    ["上午 v1 立项", "需求确认（云端大模型/智谱/系统日历）；多 LLM Provider 架构（OpenAI 兼容基类+三家）；表单式 UI + Speech 语音转写 + EventKit；手写 Xcode 工程文件"],
    ["中午 CI 搭建", "Windows 无法跑 iOS 模拟器 → GitHub Actions 方案；浏览器代操作建仓；workflow+共享 scheme；首次推送过 GCM 授权"],
    ["下午 真机侧载排障", "双产物 workflow；Sideloadly 依次遭遇：安装包写入错误 → CFBundleIdentifier 缺失 → Invalid file（英文路径/关代理/升级均无效）→ 登录 -20101/-22406 → 爱思弃用。最终两大根因：①开发者模式未开启且开关隐藏（iMazing 触发）②免费 Apple ID 必须主密码而非专用密码。真机安装成功"],
    ["晚上 v1 收尾 + v2 改版", "日历权限新键修复；v2 聊天式改版（消息流+卡片+输入栏+抽屉+批量识别+6套图标+改名 Orbit）；v2.1 七条反馈修复；相机闪退修复（NSCameraUsageDescription）"],
])

# ===== 5 踩坑记录 =====
h1("5 · 踩坑记录（12 坑全档）")
table(["#", "现象", "根因", "解法"], [
    ["1", "Windows 想跑 iOS 模拟器", "模拟器 macOS 专属", "GitHub Actions 云构建"],
    ["2", "Sideloadly 报 CFBundleIdentifier", "手写 Info.plist 缺必需字段", "GENERATE_INFOPLIST_FILE=YES"],
    ["3", "git push 连不上", "Clash 代理 git 没走", "git -c http.proxy=http://127.0.0.1:7897 push"],
    ["4", "开发者网站 302/转圈", "代理退出直连不稳/未登录先进站", "先登 appleid.apple.com 再进；必要时开代理"],
    ["5", "爱思 Get Xcode token 失败", "工具兼容性（手机号ID尤甚）", "换 Sideloadly"],
    ["6", "Sideloadly 登录 -20101", "两步验证密码形态错误", "见坑9"],
    ["7", "Sideloadly 登录 -22406", "免费ID用了App专用密码", "见坑9"],
    ["8", "Sideloadly Invalid file", "iOS16+ 开发者模式未开启（开关默认隐藏）", "iMazing 启用开发者模式→手机开启→重启"],
    ["9", "密码到底填哪个", "专用密码只对付费开发者账号有效", "手机号ID + Apple ID 主密码（弹验证码照输）"],
    ["10", "日历无弹窗/无可用日历", "缺 iOS17+ 键 NSCalendarsFullAccessUsageDescription 等", "补两组日历权限键+仅写入降级"],
    ["11", "助手回复黑框无字", "SwiftUI .background(Shape) 未填色=黑底黑字", "shape.fill(secondarySystemBackground)；用户气泡 fill(.blue)"],
    ["12", "点相机闪退", "缺 NSCameraUsageDescription 被系统杀进程", "Info.plist 补键"],
])
para("规律：权限键要一次配全（日历/相机/麦克风/语音/相库）；免费侧载三件套=开发者模式+主密码+Sideloadly；报错先查权限，再查路径，再查网络。", bold=True)

# ===== 6 安装速查卡 =====
doc.add_page_break()
h1("6 · 安装速查卡（新手机/重装/7天重签通用）")
steps = [
    ("开启开发者模式", "电脑装 iMazing 3.4+（imazing.com）→ iPhone 连线 → 『启用开发者模式』→ 手机 设置→隐私与安全性→开发者模式（拉到底）→ 打开 → 输锁屏密码 → 重启手机。只需一次，已开过的跳过。"),
    ("下载 IPA", "github.com/hanwen-FDE/orbit-scheduler → Actions → 最新绿勾 → Artifacts → 下载 ScheduleAssistant-ipa → 解压出 ScheduleAssistant.ipa 放纯英文路径（如 D:\\ipa\\）。"),
    ("Sideloadly 安装", "sideloadly.io 装 v0.60+ → iPhone 连线（信任）→ 拖入 IPA → Apple Account 填手机号ID（8613…）→ 密码填 Apple ID 主密码（不是App专用密码！弹验证码输iPhone上的6位码）→ Start。"),
    ("信任证书", "手机 设置→通用→VPN与设备管理 → 点你的 Apple ID → 信任 → 打开 Orbit。"),
    ("每7天重签", "签名到期 App 打不开时，重复第3步即可，聊天记录和日程不丢。改了代码先做第2步换新 IPA。"),
]
for i, (t, d) in enumerate(steps, 1):
    para(f"第{i}步 {t}", bold=True, color=PURPLE)
    para(d)

# ===== 7 维护手册 =====
h1("7 · 日常维护手册")
h2("7.1 改功能/修 bug 标准流程")
para("① 修改代码（Sources/ 下 Swift 文件）→ ② git add -A && git commit -m “说明” → ③ git push（需 Clash 开着）→ ④ 等 Actions 约1分钟（绿勾成功/红叉看日志）→ ⑤ 下载新 IPA 解压覆盖 D:\\ipa\\ → ⑥ Sideloadly 重装")
h2("7.2 数据说明")
para("聊天记录与日程快照存于 App 沙盒（orbit-chat.json）；日程本体在系统日历（iCloud 同步）。卸载 App 会清聊天记录，日历里的日程保留。API Key 重装覆盖不丢。")
h2("7.3 关键文件索引")
table(["文件", "职责"], [
    ["Sources/LLMProvider.swift", "多服务商抽象+Prompt+解析（换模型/调提示词改这里）"],
    ["Sources/ChatStore.swift", "消息管线：发送→识别→写日历→卡片（业务主逻辑）"],
    ["Sources/CalendarService.swift", "EventKit 增删改查/提醒/换日历"],
    ["Sources/ChatView / MessageViews", "聊天 UI 与日程卡片"],
    [".github/workflows/build.yml", "CI 双产物构建"],
    ["make_icons.py", "图标生成脚本（改配色重跑）"],
])

# ===== 8 Backlog =====
h1("8 · 后续功能规划（Backlog）")
table(["优先级", "功能", "说明"], [
    ["高", "识别准确度调优", "收集错例改 Prompt（下周三推算、跨天会议、农历）"],
    ["高", "分享菜单接入", "微信复制的会议通知→系统分享→Orbit 一键识别"],
    ["中", "剪贴板监听提示", "复制日程文本后打开 App 自动提示识别"],
    ["中", "Siri 快捷指令", "『嘿Siri记个日程』直达语音输入"],
    ["中", "日历视图", "日程页加月/周视图"],
    ["低", "云端同步", "记录/配置 iCloud 同步，换机不丢"],
    ["低", "TestFlight 分发", "若开通付费开发者账号，同事直接安装免重签"],
])

doc.save("docs/开发报告.docx")
print("docx saved")
